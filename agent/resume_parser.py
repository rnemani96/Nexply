from pathlib import Path
from docx import Document
import json
import re


BASE_DIR = Path(__file__).resolve().parent.parent
RESUME_PATH = BASE_DIR / "resume" / "master_resume.docx"
OUTPUT_DIR = BASE_DIR / "output"
PROFILE_PATH = OUTPUT_DIR / "candidate_profile.json"


def extract_resume_text():
    if not RESUME_PATH.exists():
        raise FileNotFoundError(
            f"Resume not found: {RESUME_PATH}"
        )

    document = Document(RESUME_PATH)

    sections = []

    # Normal paragraphs
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            sections.append(text)

    # Tables
    for table in document.tables:
        for row in table.rows:
            row_text = []

            for cell in row.cells:
                text = cell.text.strip()

                if text:
                    row_text.append(text)

            if row_text:
                sections.append(" | ".join(row_text))

    return "\n".join(sections)


def extract_email(text):
    pattern = r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}"

    match = re.search(pattern, text)

    return match.group(0) if match else ""


def extract_phone(text):
    pattern = r"(?<!\d)(?:\+?\d[\d\s().-]{8,}\d)(?!\d)"

    match = re.search(pattern, text)

    return match.group(0).strip() if match else ""


def find_section(text, section_names):
    lines = text.splitlines()

    start = None

    for index, line in enumerate(lines):

        normalized = line.strip().lower()

        for name in section_names:

            if normalized == name.lower():
                start = index + 1
                break

        if start is not None:
            break

    if start is None:
        return ""

    collected = []

    common_headers = {
        "summary",
        "professional summary",
        "skills",
        "technical skills",
        "experience",
        "work experience",
        "professional experience",
        "education",
        "projects",
        "certifications",
        "achievements",
    }

    for line in lines[start:]:

        normalized = line.strip().lower()

        if normalized in common_headers:
            break

        if line.strip():
            collected.append(line.strip())

    return "\n".join(collected)


def build_profile(text):

    profile = {
        "source": "master_resume.docx",

        "contact": {
            "email": extract_email(text),
            "phone": extract_phone(text),
        },

        "summary": find_section(
            text,
            [
                "summary",
                "professional summary",
                "profile",
            ],
        ),

        "skills_text": find_section(
            text,
            [
                "skills",
                "technical skills",
                "technical expertise",
            ],
        ),

        "experience_text": find_section(
            text,
            [
                "experience",
                "work experience",
                "professional experience",
            ],
        ),

        "education_text": find_section(
            text,
            [
                "education",
                "academic background",
            ],
        ),

        "projects_text": find_section(
            text,
            [
                "projects",
                "key projects",
            ],
        ),

        "certifications_text": find_section(
            text,
            [
                "certifications",
                "certificates",
            ],
        ),

        "raw_text": text,
    }

    return profile


def save_profile(profile):

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    with open(
        PROFILE_PATH,
        "w",
        encoding="utf-8"
    ) as file:

        json.dump(
            profile,
            file,
            indent=2,
            ensure_ascii=False
        )


def main():

    print("=" * 70)
    print("RAJESH AI - RESUME PARSER")
    print("=" * 70)

    print()
    print(f"Resume: {RESUME_PATH}")

    if not RESUME_PATH.exists():

        print()
        print("STATUS: WAITING FOR MASTER RESUME")
        print()
        print(
            "Place your resume at:"
        )
        print(
            RESUME_PATH
        )

        return

    print("Reading resume...")

    text = extract_resume_text()

    print(f"Characters extracted: {len(text)}")

    profile = build_profile(text)

    save_profile(profile)

    print()
    print("PROFILE CREATED")
    print("-" * 50)

    print(
        "Email:",
        profile["contact"]["email"] or "Not found"
    )

    print(
        "Phone:",
        profile["contact"]["phone"] or "Not found"
    )

    print(
        "Summary:",
        "Found" if profile["summary"] else "Not found"
    )

    print(
        "Skills:",
        "Found" if profile["skills_text"] else "Not found"
    )

    print(
        "Experience:",
        "Found" if profile["experience_text"] else "Not found"
    )

    print(
        "Education:",
        "Found" if profile["education_text"] else "Not found"
    )

    print()
    print(f"Saved: {PROFILE_PATH}")

    print()
    print("=" * 70)


if __name__ == "__main__":
    main()